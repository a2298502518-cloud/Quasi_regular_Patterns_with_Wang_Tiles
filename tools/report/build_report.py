from copy import deepcopy
from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


WORKSPACE = Path(__file__).resolve().parents[2]
TMP_DIR = Path(__file__).resolve().parent
EQUATION_DIR = TMP_DIR / "equations"
OUTPUT_PATH = WORKSPACE / "reports" / "基于QRP与Wangtile的准规则纹样生成实验报告.docx"
FLOW_PATH = TMP_DIR / "method-flow.png"
RESULTS_DIR = WORKSPACE / "output" / "wang-qrp-experiment"
MML2OMML_XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(119, 119, 119)
HEADER_GREEN = "244C32"
PALE_GREEN = "F5F6F5"
BORDER_GRAY = "D9D9D9"

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATHML_NS = "http://www.w3.org/1998/Math/MathML"
MATHML_TO_OMML = etree.XSLT(etree.parse(str(MML2OMML_XSL)))


def normalize_report_text(value):
    text = str(value)
    replacements = (
        ("Wang 瓦片（Wang tile）", "Wangtile"),
        ("Wang瓦片（Wang tile）", "Wangtile"),
        ("Wang 瓦片", "Wangtile"),
        ("Wang瓦片", "Wangtile"),
        ("Wang tile", "Wangtile"),
        ("Wang Tile", "Wangtile"),
        ("边颜色（edge colors）", "边颜色"),
        ("准规则纹样（quasi-regular pattern，QRP）", "准规则纹样，简称QRP"),
        ("颜色映射（colormap）", "颜色映射"),
        ("伪随机种子（seed）", "伪随机种子"),
        ("均方根差（root-mean-square difference，RMSD）", "均方根差，简称RMSD"),
    )
    for source, target in replacements:
        text = text.replace(source, target)
    text = re.sub(r"\bWang(?!tile)\b", "Wangtile", text)
    text = text.replace("瓦片", "Wangtile")

    cjk = r"\u3400-\u4DBF\u4E00-\u9FFF"
    token_start = r"A-Za-z\u0370-\u03FF"
    token_tail = token_start + r"0-9⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉_*+\-/=<>≤≥∈×%|."
    text = re.sub(fr"([{cjk}])\s+(?=[{token_start}])", r"\1", text)
    text = re.sub(fr"([{token_start}][{token_tail}]*)\s+(?=[{cjk}])", r"\1", text)
    text = re.sub(fr"([{cjk}])\s+(?=[(（\[])", r"\1", text)
    text = re.sub(fr"([)）\]])\s+(?=[{cjk}])", r"\1", text)
    return text


def set_font(run, name="Microsoft YaHei", size=11, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)


def set_style_font(style, name, size, bold=False, italic=False, color=BLACK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), name)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    set_font(run, size=8.5, color=GRAY)
    run._r.extend([begin, instruction, separate, text, end])


def add_body(doc, text, keep_next=False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(normalize_report_text(text))
    set_font(run)
    return paragraph


def add_heading(doc, text, level, page_break=False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break
    run = paragraph.add_run(normalize_report_text(text))
    set_font(run, size=15 if level == 1 else 12.5, bold=True)
    return paragraph


def append_math_run_properties(parent):
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), "Cambria Math")
    run_properties.append(fonts)
    for tag in ("w:sz", "w:szCs"):
        node = OxmlElement(tag)
        node.set(qn("w:val"), "20")
        run_properties.append(node)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    run_properties.append(color)
    parent.append(run_properties)


def set_math_run_properties(math_element):
    namespaces = {"m": MATH_NS}
    for math_run in math_element.xpath(".//m:r", namespaces=namespaces):
        old = math_run.find(qn("w:rPr"))
        if old is not None:
            math_run.remove(old)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{name}"), "Cambria Math")
        run_properties.append(fonts)
        for tag in ("w:sz", "w:szCs"):
            node = OxmlElement(tag)
            node.set(qn("w:val"), "20")
            run_properties.append(node)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        run_properties.append(color)
        insertion_index = 1 if math_run.find(qn("m:rPr")) is not None else 0
        math_run.insert(insertion_index, run_properties)


def load_native_equation(equation_id):
    mathml_path = EQUATION_DIR / f"{equation_id}.mml"
    mathml = etree.parse(str(mathml_path))
    mathml_namespaces = {"mml": MATHML_NS}
    for accent in mathml.xpath(
        "//mml:mover[mml:mo='^' or mml:mo='¯' or mml:mo='~']",
        namespaces=mathml_namespaces,
    ):
        accent.set("accent", "true")
    transformed = MATHML_TO_OMML(mathml)
    root = transformed.getroot()
    if root.tag == qn("m:oMathPara"):
        root = root.find(qn("m:oMath"))
    if root is None or root.tag != qn("m:oMath"):
        raise ValueError(f"Unable to convert {equation_id} to native Word math")
    equation = deepcopy(root)
    source_table = mathml.getroot().find(f"{{{MATHML_NS}}}mtable")
    matrix = equation.find(qn("m:m"))
    if source_table is not None and matrix is not None:
        spacings = (source_table.get("columnspacing") or "").split()
        equation_array = OxmlElement("m:eqArr")
        for row in matrix.findall(qn("m:mr")):
            line = OxmlElement("m:e")
            cells = row.findall(qn("m:e"))
            for index, cell in enumerate(cells):
                has_text = any(
                    text_node.text or ""
                    for text_node in cell.xpath(".//m:t", namespaces={"m": MATH_NS})
                )
                if has_text:
                    for child in list(cell):
                        line.append(deepcopy(child))
                if index < len(cells) - 1 and index < len(spacings) and spacings[index] != "0em":
                    spacer = OxmlElement("m:r")
                    text_node = OxmlElement("m:t")
                    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    text_node.text = "\u2003\u2003"
                    spacer.append(text_node)
                    line.append(spacer)
            equation_array.append(line)
        equation.replace(matrix, equation_array)
    set_math_run_properties(equation)
    return parse_xml(etree.tostring(equation))


def add_equation(doc, equation_id, number):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(0.65), Inches(5.55), Inches(0.65))
    for index, width in enumerate(widths):
        table.columns[index].width = width
        cell = table.cell(0, index)
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=35, start=0, bottom=55, end=0)
    remove_table_borders(table)
    center = table.cell(0, 1).paragraphs[0]
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center.paragraph_format.space_after = Pt(0)
    center.paragraph_format.first_line_indent = Pt(0)
    center.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    center.paragraph_format.line_spacing = 1.0
    append_math_run_properties(center._p.get_or_add_pPr())
    center._p.append(load_native_equation(equation_id))
    number_paragraph = table.cell(0, 2).paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    number_paragraph.paragraph_format.space_after = Pt(0)
    number_paragraph.paragraph_format.first_line_indent = Pt(0)
    number_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    number_paragraph.paragraph_format.line_spacing = 1.0
    run = number_paragraph.add_run(f"({number})")
    set_font(run, name="Cambria Math", size=10)
    row_pr = table.rows[0]._tr.get_or_add_trPr()
    row_pr.append(OxmlElement("w:cantSplit"))
    return table


def create_flow_figure():
    width, height = 1800, 300
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    bold_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 31)
    body_font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 22)
    labels = [
        ("Wang 边标签", "合法铺砌"),
        ("C¹ 权重场", "Hermite 延拓"),
        ("QRP 相位调制", "固定 Wang 参数"),
        ("粗尺度结构", "QRP 参数控制"),
        ("多尺度细节", "中细尺度合成"),
        ("颜色映射", "固定显示范围"),
    ]
    margin, box_width, box_height, gap = 45, 240, 118, 54
    y0 = 76
    for index, (title, subtitle) in enumerate(labels):
        x0 = margin + index * (box_width + gap)
        x1 = x0 + box_width
        y1 = y0 + box_height
        draw.rounded_rectangle(
            (x0, y0, x1, y1), radius=14, fill="#F5F6F5", outline="#244C32", width=3
        )
        draw.text(
            ((x0 + x1) / 2, y0 + 39),
            normalize_report_text(title),
            font=bold_font,
            fill="#111111",
            anchor="mm",
        )
        draw.text(
            ((x0 + x1) / 2, y0 + 84),
            normalize_report_text(subtitle),
            font=body_font,
            fill="#444444",
            anchor="mm",
        )
        if index + 1 < len(labels):
            arrow_start = x1 + 10
            arrow_end = x1 + gap - 10
            y = y0 + box_height / 2
            draw.line((arrow_start, y, arrow_end, y), fill="#244C32", width=4)
            draw.polygon(
                ((arrow_end, y), (arrow_end - 15, y - 9), (arrow_end - 15, y + 9)),
                fill="#244C32",
            )
    image.save(FLOW_PATH, dpi=(300, 300))


def add_figure(doc, image_path, width, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(normalize_report_text(caption))
    set_font(caption_run, size=9.5)


def add_role_table(doc):
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    set_font(caption.add_run(normalize_report_text("表 1 模型组成与参数职责")), size=9.5)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.35), Inches(1.55), Inches(3.55))
    for index, width in enumerate(widths):
        table.columns[index].width = width
    headers = ("组成", "实验设置", "主要作用")
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        cell.width = widths[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade_cell(cell, HEADER_GREEN)
        set_cell_margins(cell, top=120, start=100, bottom=120, end=100)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(
            paragraph.add_run(normalize_report_text(text)),
            size=9.5,
            bold=True,
            color=RGBColor(255, 255, 255),
        )
    set_repeat_table_header(table.rows[0])

    rows = (
        ("Wang 铺砌", "固定", "规定边标签的合法排列和共享边匹配关系"),
        ("连续权重场", "固定", "将离散边标签延拓为全局 C¹ 连续的相位调制量"),
        ("QRP 标量场", "主要调节对象", "通过 q、κ、A₀、B₀ 控制粗尺度空间结构与尺度"),
        ("多尺度细节与颜色映射", "结构对照中固定", "补充中细尺度细节并统一显示条件"),
    )
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                shade_cell(cell, PALE_GREEN)
        for column, text in enumerate(values):
            cell = cells[column]
            cell.width = widths[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=75, start=90, bottom=75, end=90)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column < 2 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(13.5)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            set_font(paragraph.add_run(normalize_report_text(text)), size=9.5)
        row_pr = table.rows[-1]._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_pr.append(cannot_split)
    set_table_borders(table)


def add_data_table(doc, number, title, headers, rows, widths, alignments=None):
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    set_font(caption.add_run(normalize_report_text(f"表 {number} {title}")), size=9.5)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    column_widths = tuple(Inches(value) for value in widths)
    for index, width in enumerate(column_widths):
        table.columns[index].width = width
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        cell.width = column_widths[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade_cell(cell, HEADER_GREEN)
        set_cell_margins(cell, top=80, start=80, bottom=80, end=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(
            paragraph.add_run(normalize_report_text(text)),
            size=9.2,
            bold=True,
            color=RGBColor(255, 255, 255),
        )
    set_repeat_table_header(table.rows[0])

    if alignments is None:
        alignments = [WD_ALIGN_PARAGRAPH.CENTER] * len(headers)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                shade_cell(cell, PALE_GREEN)
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = column_widths[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=45, start=78, bottom=45, end=78)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignments[column]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(13)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            set_font(paragraph.add_run(normalize_report_text(value)), size=9.2)
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    set_table_borders(table)
    return table


def remove_paragraph_borders(paragraph_properties):
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def configure_document(doc):
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.68)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Microsoft YaHei", 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.3056)
    normal.paragraph_format.line_spacing = Pt(17.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    set_style_font(title, "Microsoft YaHei", 25, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = Pt(32)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title._element.get_or_add_pPr())

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "Microsoft YaHei", 12, italic=True)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True

    for level, size, line in ((1, 15, 22), (2, 12.5, 18)):
        heading = styles[f"Heading {level}"]
        set_style_font(heading, "Microsoft YaHei", size, bold=True)
        heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.first_line_indent = Pt(0)
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(7)
        heading.paragraph_format.line_spacing = Pt(line)
        heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True

    caption = styles["Caption"]
    set_style_font(caption, "Microsoft YaHei", 9.5)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.line_spacing = Pt(13)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    add_page_number(footer_paragraph)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build_report():
    create_flow_figure()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = normalize_report_text(
        "基于 QRP 与 Wang 瓦片的准规则纹样生成实验报告"
    )
    doc.core_properties.subject = normalize_report_text(
        "Wang 边标签诱导的连续二阶相位调制方法"
    )

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title._p.get_or_add_pPr())
    set_font(
        title.add_run(normalize_report_text("基于 QRP 与 Wang 瓦片的准规则纹样生成实验报告")),
        size=25,
        bold=True,
    )
    subtitle = doc.add_paragraph(style="Subtitle")
    set_font(
        subtitle.add_run(normalize_report_text("Wang 边标签诱导的连续二阶相位调制方法")),
        size=12,
        italic=True,
    )

    add_heading(doc, "1 实验目的", 1)
    add_body(doc, "规则单元在网格中的反复复制容易形成明显的平移周期，较大范围内的结构变化因而受到限制。Wang 瓦片（Wang tile）在边上配置离散匹配符号，经典文献称其为边颜色（edge colors）；本文称为边标签，以免与最终的颜色映射混淆。准规则纹样（quasi-regular pattern，QRP）由多个方向上的共振分量共同构成，可形成具有统一生成规律和复杂局部变化的纹样。二者结合后，边标签负责空间组织，共振参数负责纹理结构。")
    add_body(doc, "本实验研究一种 QRP 与 Wang 瓦片耦合的程序化纹理模型。纹理在全局坐标中由同一个解析场求值，流程不预先生成瓦片图像，也不为每一种瓦片类型单独存储纹理。Wang 网格、边界数据函数和耦合强度在结构对照中保持固定；粗尺度空间结构主要由 QRP 的共振方向数、空间波数和方向间相位关系控制。这样的参数分工使不同结果之间的变化来源可以直接追溯。")
    add_body(doc, "方法首先把离散边标签延拓为全局 C¹ 连续的权重场，使共享边两侧具有一致的函数值及全部一阶偏导数。该权重随后进入 QRP 相位，参与局部结构的连续调制。粗尺度结构确定后，模型以统一的中、细尺度场补充局部层次。整个过程不使用逐瓦片随机相位、随机旋转或附加噪声；给定模型参数和铺砌的伪随机种子后，可以重复得到相同结果。")
    add_body(doc, "本阶段的目标是建立可验证的方法基础：在 Wang 网格和耦合参数固定的条件下，同一公式应当能够通过 QRP 参数产生清楚的结构差异；共享边上的函数值及一阶偏导数应保持一致。实验重点检验 QRP 参数的结构控制能力，并以数值测试确认连续权重场和参数化 QRP 场的正确性。")

    add_heading(doc, "2 方法", 1, page_break=True)
    add_heading(doc, "2.1 整体框架", 2)
    add_body(doc, "方法由 Wang 铺砌、连续权重构造、参数化 QRP 标量场、多尺度纹理合成和颜色映射五个环节组成。图 1 给出了从离散边标签到最终标量纹理的计算顺序。所有环节使用同一套全局坐标；瓦片只提供边标签及局部坐标，不重置 QRP 的坐标原点。", keep_next=True)
    add_figure(doc, FLOW_PATH, 6.65, "图 1 方法的整体计算流程")
    add_body(doc, "Wang 配置在本阶段确定后保持不变，QRP 参数承担纹理结构调控。中、细尺度参数和颜色映射（colormap）在结构对照中同样固定，使结果差异能够归因于 QRP。各部分的职责见表 1。", keep_next=True)
    add_role_table(doc)

    add_heading(doc, "2.2 Wang 铺砌与边标签", 2)
    add_body(doc, "设全局坐标为 x，位于整数网格位置 g=(i,j) 的单位正方形瓦片使用局部坐标 ξ=(u,v)。瓦片不旋转也不镜像，其有序边标签由南、北、西、东四个标签组成，标签编号取自 K 个有限类别。", keep_next=True)
    add_equation(doc, "eq01", 1)
    add_body(doc, "南、北边的参数方向统一规定为从左到右，西、东边统一规定为从下到上。水平方向和竖直方向的合法邻接条件分别为", keep_next=True)
    add_equation(doc, "eq02", 2)
    add_body(doc, "边标签本身是离散符号，连续模型通过标签到边界数据的确定性映射进入后续计算。统一的边参数方向同时用于边界值和跨边界方向导数：南、北边使用沿全局 +y 方向的 ∂W/∂v，西、东边使用沿全局 +x 方向的 ∂W/∂u。该导数不按瓦片外法向定义，因此共享边两侧可以直接复用同一组一阶数据而无需改变符号。")
    add_body(doc, "Wang 网格的伪随机种子（seed）只决定合法标签在平面中的排列，不进入连续纹理公式。结构参数实验固定网格尺寸、标签数量、随机种子和边界数据函数。QRP 参数改变时，Wang 排列保持一致；观察到的粗尺度结构变化因此不会与铺砌布局变化混合。")

    add_heading(doc, "2.3 由 Wang 边标签构造的全局 C¹ 连续权重场", 2)
    add_body(doc, "离散边标签通过确定性的标量场 W(x) 进入 QRP。W(x) 用于提供连续相位调制量，不直接作为最终纹理。对于标签 c，定义标签角和端点平坦的气泡函数", keep_next=True)
    add_equation(doc, "eq03", 3)
    add_body(doc, "标签对应的边界值函数和跨边界方向导数函数写为", keep_next=True)
    add_equation(doc, "eq04", 4)
    add_body(doc, "其中，w* 为基准权重，ρ 控制不同标签的边界值差异，σ 控制跨边界方向一阶导数的差异。气泡函数及其导数在端点处均为零，因此全部标签在瓦片角点共享相同的一阶数据。", keep_next=True)
    add_equation(doc, "eq05", 5)
    add_body(doc, "四条边的数据使用三次 Hermite 基延拓到瓦片内部。基函数为", keep_next=True)
    add_equation(doc, "eq06", 6)
    add_body(doc, "从各边界值中扣除基准值后，南北边和东西边的内部贡献分别定义为", keep_next=True)
    add_equation(doc, "eq07", 7)
    add_equation(doc, "eq08", 8)
    add_body(doc, "瓦片内部的统一权重函数为", keep_next=True)
    add_equation(doc, "eq09", 9)
    add_body(doc, "Pₓ 在南北边上连同一阶导数一起消失，Pᵧ 在东西边上具有相同性质。两组延拓可以同时插值四条边的函数值、切向导数和跨边界方向导数，不需要逐瓦片求解或角点修正。Hermite 基的范围进一步给出以下值域充分条件。", keep_next=True)
    add_equation(doc, "eq10", 10)
    add_body(doc, "当式（10）成立时，全部合法边标签组合均满足 0≤W(x)≤1。相邻瓦片在共享边上使用同一标签、同一参数方向和同一跨边界方向导数约定，因此 W 及其梯度向量逐点一致，分片构成的权重场在整个 Wang 网格上达到全局 C¹ 连续。")

    add_heading(doc, "2.4 参数化 QRP 标量场", 2)
    add_body(doc, "本实验采用基本 QRP 模型的整数共振方向数 q 归一化特例，并以空间波数 κ 调节纹理尺度。第 r 个方向由角度 θr 和单位方向向量 er 表示。", keep_next=True)
    add_equation(doc, "eq11", 11)
    add_body(doc, "不含附加相位时，归一化基本 QRP 场写为", keep_next=True)
    add_equation(doc, "eq12", 12)
    add_body(doc, "参数 q 决定参与干涉的方向数量，空间波数 κ>0 控制空间尺度。余弦分量采用按 q 归一化的等权平均，使不同 q 下的幅值具有可比性，并保证场值位于 [-1,1]。各方向的一阶角谐波相位等价于一次共同坐标平移：", keep_next=True)
    add_equation(doc, "eq13", 13)
    add_body(doc, "这类参数主要移动已有结构。以下二阶角谐波相位参数化是本实验在基本 QRP 模型上的扩展，用于调节等角分布的共振方向集之间的相对相位；同时将 Wang 权重中心化为 U(x)。", keep_next=True)
    add_equation(doc, "eq14", 14)
    add_body(doc, "其中，式（14）中的 RW=2|ρ|+|σ|/2 是权重相对基准值的解析偏差上界；RW>0 时有 |U(x)|≤1，RW=0 时定义 U(x)=0。两个连续相位系数场定义为", keep_next=True)
    add_equation(doc, "eq15", 15)
    add_body(doc, "A₀ 和 B₀ 是全局二阶角谐波相位调制系数；式（15）中的 λA 和 λB 是固定的 Wang 权重与相位耦合系数。记参数集合 Θ=(q,κ,A₀,B₀,λA,λB)，最终的粗尺度场记为 CΘ(x)，下文简记为 C(x)。", keep_next=True)
    add_equation(doc, "eq16", 16)
    add_body(doc, "当 A₀、B₀、λA 和 λB 均为零时，式（16）退化为归一化基本 QRP 场。A₀ 和 B₀ 改变同一共振方向集内部各方向的相位关系，从而调节干涉轮廓；q 和 κ 分别调节方向组成与结构尺度。本阶段固定 Wang 网格、w*、ρ、σ、λA 和 λB，主要通过 q、κ、A₀ 和 B₀ 控制纹理。具体参数能否形成稳定且可辨认的粗尺度结构差异，将在实验结果中比较。")

    add_heading(doc, "2.5 多尺度纹理合成", 2)
    add_body(doc, "粗尺度场 C(x) 决定主要轮廓。模型另外计算中尺度场 Dm(x) 和细尺度场 Df(x)，并使用门控函数将二者加入粗结构。", keep_next=True)
    add_equation(doc, "eq17", 17)
    add_body(doc, "βm 和 βf 分别控制中尺度与细尺度成分的强度。门控项 1-C² 在粗场接近正负极值时趋近于零，在过渡区域取得较大值，因此细节主要进入轮廓附近，不会大面积覆盖稳定的极值区域。")
    add_body(doc, "Dm 和 Df 均按式（16）的确定性 QRP 形式在全局坐标中计算，分别使用固定的中尺度参数组和细尺度参数组，并将二者的 Wang 权重与相位耦合系数设为零。它们采用高于粗尺度场的空间波数，且在所有结构对照中保持 Dm、Df、βm 和 βf 完全一致。合成过程不执行逐图最小值－最大值归一化，所有结果使用同一标量区间和颜色映射，从而保留参数组之间的相对幅度关系。")

    add_heading(doc, "2.6 连续性与值域", 2)
    add_body(doc, "记第 r 个共振分量的总相位为 χr。其空间梯度为", keep_next=True)
    add_equation(doc, "eq18", 18)
    add_body(doc, "由链式法则可得粗尺度场的解析梯度", keep_next=True)
    add_equation(doc, "eq19", 19)
    add_body(doc, "权重偏差上界 RW 为常数时，相位系数场的梯度直接由 Wang 权重场给出。", keep_next=True)
    add_equation(doc, "eq20", 20)
    add_body(doc, "W(x) 在合法 Wang 网格上全局 C¹ 连续，因此 A(x)、B(x) 和式（16）中的有限余弦和 CΘ(x) 也保持全局 C¹ 连续。共享边两侧的场值和梯度向量由同一解析表达式确定，连续性不依赖像素级模糊或边界后处理。")
    add_body(doc, "粗尺度场是余弦分量的平均值，其幅值不超过 1。层级参数满足下列充分条件时，合成结果同样保持在 [-1,1] 内。", keep_next=True)
    add_equation(doc, "eq21", 21)
    add_body(doc, "共享边采用一致的参数方向，各级 QRP 场均在同一全局坐标中求值。连续颜色映射随后把标量场映射为颜色，因而最终图像能够继承共享边上的连续性。")

    add_heading(doc, "3 实验设计", 1, page_break=True)
    add_heading(doc, "3.1 实验内容", 2)
    add_body(doc, "实验围绕两项内容展开。第一项考察固定 Wang 铺砌后，QRP 全局相位参数能否在同一解析模型中生成具有明显差异的粗尺度纹理结构。第二项检查连续权重场、参数化 QRP 场和共享边一阶数据的数值正确性。")
    add_body(doc, "粗尺度场是本轮视觉比较的主要对象。各组采用相同画布、相同 Wang 网格、相同颜色映射和固定显示范围，避免铺砌布局或逐图归一化对观察结果产生干扰。所有连续纹理均由确定性公式生成，不加入随机噪声、逐瓦片随机相位、随机旋转和事后模糊。")

    add_heading(doc, "3.2 固定条件与参数组", 2)
    add_body(doc, "所有实验使用 20×20 的 Wang 网格，即画面包含 400 个瓦片实例；每个瓦片以 80×80 像素采样，单幅结果分辨率为 1600×1600。边标签类别数固定为 5，铺砌的伪随机种子固定为 0x4d595df4d0f33173。该种子只用于生成可复现的合法标签排列，不参与连续标量场的数值求值。方法预先生成的瓦片图像数量为 0，主要固定参数见表 2。", keep_next=True)
    add_data_table(
        doc,
        2,
        "实验固定参数",
        ("参数类别", "参数设置", "实验作用"),
        (
            ("Wang 网格", "20×20；5 种边标签", "固定 400 个瓦片实例的布局与邻接约束"),
            ("采样", "80×80 像素/瓦片", "得到 1600×1600 输出"),
            ("权重场", "K=5，w*=0.5，ρ=0.10，σ=0.40", "生成连续局部调制量"),
            ("粗尺度 QRP", "q=7，κ=3.15", "承载主要纹理结构"),
            ("Wang 相位耦合", "λA=1.15，λB=-0.80", "固定权重与相位的耦合强度"),
        ),
        (1.32, 2.55, 2.58),
        (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT),
    )
    add_body(doc, "由式（10）可得本组参数的权重偏差上界 RW=0.40，因此 W(x)∈[0.1,0.9]，中心化调制量 U(x)∈[-1,1]。结构参数实验只改变粗尺度场的全局二阶角谐波相位系数 (A₀,B₀)，具体设置见表 3。", keep_next=True)
    add_data_table(
        doc,
        3,
        "QRP 粗尺度结构参数组",
        ("组别", "A₀", "B₀", "参数设置"),
        (
            ("P0", "0.00", "0.00", "零全局相位偏置（保留 Wang 调制）"),
            ("P1", "1.65", "0.00", "仅 A₀ 非零"),
            ("P2", "3.00", "0.85", "A₀、B₀ 同号"),
            ("P3", "2.10", "-2.45", "A₀、B₀ 异号"),
        ),
        (0.82, 1.05, 1.05, 3.53),
        (WD_ALIGN_PARAGRAPH.CENTER,) * 4,
    )

    add_heading(doc, "3.3 QRP 结构参数实验", 2)
    add_body(doc, "程序分别计算 P0、P1、P2 和 P3 的粗尺度场。四组结果采用同一 Wang 网格、耦合系数、采样分辨率和颜色映射，仅改变 (A₀,B₀)。视觉比较主要观察明暗区域的空间分布、轮廓方向、局部聚集形式和连接关系。")
    add_body(doc, "为辅助判断参数变化是否已经作用于标量场，实验在相同像素坐标上计算任意两组粗尺度场的均方根差（root-mean-square difference，RMSD）。", keep_next=True)
    add_equation(doc, "eq22", 22)
    add_body(doc, "其中，N 为像素总数，式（22）中的 Ci(xp) 和 Cj(xp) 分别表示第 i、j 组粗尺度场在采样点 xp 的值。实验还以 0 为统一阈值，统计正负值区域分类不同的像素比例。", keep_next=True)
    add_equation(doc, "eq23", 23)
    add_body(doc, "式（22）的 RMSD 描述标量场的总体差异，式（23）描述固定阈值下的符号不一致率。两项指标与图像中的轮廓变化共同分析，用于判断同一公式下的参数控制能力。")

    add_heading(doc, "3.4 连续性与数值验证", 2)
    add_body(doc, "本轮 σ=0.40 的参数满足式（10）给出的值域充分条件。实现层面的独立单元测试另枚举 5 种标签构成的全部 5⁴=625 种有序边标签签名；该数量表示可能的标签组合，不是画面中的瓦片实例数。测试检查权重构造以及兼容共享边上的函数值和梯度向量。参数化 QRP 场测试零相位退化、解析梯度与中心差分的一致性及场值范围；多尺度合成测试参数合法性、解析梯度和值域条件。当前共有 9 项自动化数值测试，覆盖数学基础、基本 QRP 场、Wang 权重场、参数化 QRP 与 Wang 瓦片耦合场、多尺度合成和完整生成流程，9 项测试均已通过。")

    add_heading(doc, "4 实验结果与分析", 1, page_break=True)
    add_body(doc, "图 2 给出了四组粗尺度场。四组实验的 Wang 网格、边界数据函数、耦合系数、共振方向数 q、空间波数 κ、颜色映射和显示范围均保持一致，唯一变量是式（15）中的全局二阶角谐波相位系数 (A₀,B₀)。这些参数随后进入式（16），调节同一共振方向集中各方向之间的相对相位。")
    add_body(doc, "左上 P0 采用 (A₀,B₀)=(0,0)，表示零全局相位偏置，但仍保留式（15）中由 λA U(x) 与 λB U(x) 给出的 Wang 调制；右上 P1 仅令 A₀=1.65；左下 P2 采用同号组合 (3.00,0.85)；右下 P3 采用异号组合 (2.10,-2.45)。P0 形成较均衡的团块与环状组合，P1 的主要区域呈现更明显的定向伸展，P2 中若干轮廓沿斜向连接，P3 则形成另一组尺度更宽、连接走向不同的明暗区域。", keep_next=True)
    add_figure(
        doc,
        RESULTS_DIR / "A_coarse_parameter_family.png",
        6.05,
        "图 2 QRP 参数族的粗尺度场（左上 P0，右上 P1，左下 P2，右下 P3）",
    )
    add_body(doc, "粗尺度场的两两比较结果见表 4。六组 RMSD 均不低于 0.276631，零阈值符号不一致率均超过 34%。P2 与 P3 的差异最大，RMSD 为 0.447371，符号不一致率为 61.16%；P1 与 P2 的差异相对较小，但仍有 34.20% 的采样位置在两组结果中具有不同符号。", keep_next=True)
    add_data_table(
        doc,
        4,
        "QRP 粗尺度场的两两差异",
        ("参数组", "RMSD", "零阈值符号不一致率"),
        (
            ("P0–P1", "0.291436", "36.69%"),
            ("P0–P2", "0.436893", "59.82%"),
            ("P0–P3", "0.438151", "60.49%"),
            ("P1–P2", "0.276631", "34.20%"),
            ("P1–P3", "0.389104", "51.56%"),
            ("P2–P3", "0.447371", "61.16%"),
        ),
        (2.15, 2.15, 2.15),
        (WD_ALIGN_PARAGRAPH.CENTER,) * 3,
    )
    add_body(doc, "这些数据说明 (A₀,B₀) 已经改变标量场本身，所得差异超出了调色和局部亮度调整的范围。结合图 2 可以确认，统一公式能够通过少量连续参数生成多组可辨认的粗尺度结构。")

    add_heading(doc, "5 结论", 1)
    add_body(doc, "本实验建立了一种 QRP 与 Wang 瓦片耦合的程序化纹理方法。方法将 Wang 边标签延拓为全局 C¹ 连续权重场，再以该权重场调制 QRP 共振方向集中的二阶角谐波相位。纹理始终在全局坐标中由同一解析式求值，不需要预先生成瓦片图像，也不使用逐瓦片随机噪声。")
    add_body(doc, "在 Wang 网格、边界数据函数和耦合强度固定的条件下，四组全局 QRP 相位参数生成了可区分的粗尺度纹理。两两 RMSD 为 0.276631 至 0.447371，零阈值符号不一致率为 34.20% 至 61.16%，说明少量连续参数已经能够改变统一解析场中的主要纹理结构。")
    add_body(doc, "本轮权重参数满足式（10）给出的值域充分条件，独立单元测试枚举了全部 625 种有序边标签签名，当前 9 项自动化数值测试均已通过。现阶段的实验表明：在固定 Wang 框架下，QRP 参数具有可验证的结构调控能力，共享边的 C¹ 连续性也得到解析推导与数值测试的共同支持。后续可在同一框架下继续考察共振方向数 q、空间波数 κ 以及更广范围的 (A₀,B₀)，逐步扩大可生成纹理的结构类型。")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()
